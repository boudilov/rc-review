# Модалка «Помесячно по направлениям» на страницах мастерских «топ-5 + направления»:
# кнопка рядом с названием мастерской, в модалке — помесячные столбики часов SELF и CLASS
# за 2024/25 и 2025/26 с разбивкой на группы ВО / ДПО / Uteens / СПО.
#   python tools/directions_modal.py [page ...]
# Повторный запуск заменяет ранее вставленное (кнопку, модалку, стили).
import re, sys, glob, collections
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
# страница «топ-5 + направления»: префикс файлов в Stats 2026
PAGES = {'28': '01', '30': '02', '32': '03', '34': '04', '36': '05', '38': '06', '40': '07', '42': '08', '51': '09', '48': '10'}
MONTHS = [f'{y + (m < 8)}-{m:02d}' for y in (2024, 2025) for m in (*range(8, 13), *range(1, 8))]  # авг 2024 … июл 2026
MON = ['АВГ', 'СЕН', 'ОКТ', 'НОЯ', 'ДЕК', 'ЯНВ', 'ФЕВ', 'МАР', 'АПР', 'МАЙ', 'ИЮН', 'ИЮЛ']
MON_FULL = ['август', 'сентябрь', 'октябрь', 'ноябрь', 'декабрь', 'январь', 'февраль', 'март', 'апрель', 'май', 'июнь', 'июль']

# СПО UCampus не собирал — часы вписываются вручную:
# SPO[(префикс мастерской, 'SELF' | 'CLASS')] = {'2025-09': 12.5, ...}
SPO = {}

GROUPS = [  # ключ, подпись, заливка
    ('vo', 'ВО', 'rgba(255,255,255,0.92)'),
    ('dpo', 'ДПО', 'rgba(255,255,255,0.42)'),
    ('teens', 'Uteens', 'url(#{hatch})'),
    ('spo', 'СПО', 'rgba(255,255,255,0.18)'),
]
GROUP_NOTE = {
    'vo': 'бакалавриат, магистратура, специалитет',
    'dpo': 'ДПО, дополнительное образование, подготовительные курсы',
    'teens': 'Школа подростков',
    'spo': 'данные вносятся вручную',
}


def group_of(school, direction, program):
    if school == 'Школа подростков' or re.search(r'teens', str(program), re.I):
        return 'teens'
    if str(direction).startswith('Высшее') or direction == 'Бакалавриат':
        return 'vo'
    return 'dpo'


def load(prefix, kind):
    f = glob.glob(str(ROOT / 'Stats 2026' / f'{prefix}_*_{kind}_*.xlsx'))[0]
    ws = openpyxl.load_workbook(f, data_only=True).active
    hdr = [c.value for c in ws[1]]
    ix = {k: hdr.index(k) for k in ('Школа', 'Направление', 'Программа', 'Период', 'Часы')}
    data = collections.defaultdict(lambda: collections.Counter())
    progs = collections.defaultdict(lambda: collections.Counter())  # месяц → (программа, группа) → часы
    for r in ws.iter_rows(min_row=2, values_only=True):
        m = r[ix['Период']]
        if m in MONTHS:
            g = group_of(r[ix['Школа']], r[ix['Направление']], r[ix['Программа']])
            data[m][g] += r[ix['Часы']] or 0
            progs[m][(str(r[ix['Программа']] or '—').strip(), g)] += r[ix['Часы']] or 0
    for m, v in SPO.get((prefix, kind), {}).items():
        if m in MONTHS: data[m]['spo'] += v
    return data, progs


def hours(v):  # 1 234,5
    return f'{v:,.1f}'.replace(',', '&nbsp;').replace('.', ',')


def esc(s):
    return s.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;').replace('"', '&quot;')


def detail_id(page, kind, i):
    return f'ws{page}-dir-{kind.lower()}-{i:02d}'


def details(page, kind, data, progs):
    # карточка месяца: топ-10 программ; показывается по :target, возврат — ссылкой на модалку
    label = dict((k, l) for k, l, _ in GROUPS)
    out = []
    for i, m in enumerate(MONTHS):
        top = progs[m].most_common(10)
        total = sum(data[m].values())
        head = (f'<a href="#ws{page}-dir" class="wsdir-back">&larr; к графику</a>'
                f'<div class="wsdir-dh">{MON_FULL[i % 12].capitalize()} {m[:4]} · {kind}</div>'
                f'<div class="wsdir-dsub">всего {num(total)}&nbsp;ч · топ-{len(top)} программ из {len(progs[m])}</div>')
        if not top:
            rows = '<div class="wsdir-dempty">В этом месяце часов нет</div>'
        else:
            vmax = top[0][1] or 1
            rows = ''.join(
                f'<div class="wsdir-drow"><span class="wsdir-dn">{n + 1}</span>'
                f'<span class="wsdir-dp" title="{esc(p)}">{esc(p)}</span>'
                f'<span class="wsdir-dg wsdir-dg-{g}">{label[g]}</span>'
                f'<span class="wsdir-dbar"><i style="width:{100 * h / vmax:.1f}%"></i></span>'
                f'<span class="wsdir-dv">{hours(h)}&nbsp;ч</span></div>'
                for n, ((p, g), h) in enumerate(top))
        out.append(f'<div id="{detail_id(page, kind, i)}" class="wsdir-detail"><div class="wsdir-dcard">{head}{rows}</div></div>')
    return ''.join(out)


def num(v):
    return f'{round(v):,}'.replace(',', '&nbsp;')


def nice_step(vm):
    for step in (5, 10, 20, 25, 50, 100, 200, 250, 500, 1000):
        if vm / step <= 4: return step
    return 2000


def chart(page, kind, data):
    W, H, TOP, BOT = 960, 160, 6, 140  # viewBox; столбики от BOT вверх
    slot = W / 24
    bw = slot * 0.62
    totals = [sum(data[m].values()) for m in MONTHS]
    vm = max(totals) or 1
    step = nice_step(vm)
    ymax = step * -(-vm // step)
    sy = (BOT - TOP) / ymax
    hatch = f'ws{page}dir-hatch'
    out = [f'<svg class="wsdir-svg" viewBox="0 0 {W} {H}" preserveAspectRatio="none" width="100%" height="{H}">',
           f'<defs><pattern id="{hatch}" width="5" height="5" patternUnits="userSpaceOnUse" patternTransform="rotate(45)">'
           f'<rect width="5" height="5" fill="rgba(255,255,255,0.12)"/><rect width="2" height="5" fill="rgba(255,255,255,0.85)"/></pattern></defs>']
    for v in range(step, int(ymax) + 1, step):
        y = BOT - v * sy
        out.append(f'<line x1="0" y1="{y:.1f}" x2="{W}" y2="{y:.1f}" stroke="rgba(255,255,255,0.08)" stroke-dasharray="3 4" vector-effect="non-scaling-stroke"/>')
    out.append(f'<line x1="0" y1="{BOT}" x2="{W}" y2="{BOT}" stroke="rgba(255,255,255,0.18)" vector-effect="non-scaling-stroke"/>')
    out.append(f'<line x1="{W / 2}" y1="0" x2="{W / 2}" y2="{H}" stroke="rgba(255,255,255,0.22)" vector-effect="non-scaling-stroke"/>')
    for i, m in enumerate(MONTHS):
        x = i * slot + (slot - bw) / 2
        y = BOT
        parts = []
        for key, label, fill in GROUPS:
            v = data[m][key]
            if v <= 0: continue
            h = v * sy
            y -= h
            parts.append(f'<rect x="{x:.1f}" y="{y:.1f}" width="{bw:.1f}" height="{h:.1f}" fill="{fill.format(hatch=hatch)}"/>')
        tip = f'{MON_FULL[i % 12]} {m[:4]} — всего {round(totals[i])} ч' + ''.join(
            f'\n{label}: {round(data[m][key])} ч' for key, label, _ in GROUPS if data[m][key] > 0)
        tip += '\nклик — топ-10 программ месяца'
        # прозрачная колонка на весь слот — подсказка при наведении и клик на любом месяце
        out.append(f'<a href="#{detail_id(page, kind, i)}" class="wsdir-col"><title>{tip}</title><rect x="{i * slot:.1f}" y="0" width="{slot:.1f}" height="{BOT}" fill="transparent"/>{"".join(parts)}</a>')
    out.append('</svg>')
    ticks = ''.join(f'<span class="wsdir-tick" style="top:{BOT - v * sy:.1f}px">{num(v)}</span>' for v in range(0, int(ymax) + 1, step))
    months = ''.join(f'<span>{MON[i % 12]}</span>' for i in range(24))
    return (f'<div class="wsdir-plot">{"".join(out)}{ticks}</div>'
            f'<div class="wsdir-months">{months}</div>'
            f'<div class="wsdir-years"><span>2024/25</span><span>2025/26</span></div>')


def legend(data):
    y1, y2 = MONTHS[:12], MONTHS[12:]
    t1 = {k: sum(data[m][k] for m in y1) for k, *_ in GROUPS}
    t2 = {k: sum(data[m][k] for m in y2) for k, *_ in GROUPS}
    s1, s2 = sum(t1.values()) or 1, sum(t2.values()) or 1
    # таблица: столбцы — группы, строки — учебные годы (чтобы не читалось как «левая колонка = 2024/25»)
    head = ''.join(f'<th><i class="wsdir-sw wsdir-sw-{key}"></i><b>{label}</b></th>'
                   for key, label, _ in GROUPS)
    rows = ''
    for cls, year, tot, s in (('y1', '2024/25', t1, s1), ('y2', '2025/26', t2, s2)):
        cells = ''
        for key, *_ in GROUPS:
            if key == 'spo' and t1[key] == 0 and t2[key] == 0:
                cells += '<td class="wsdir-na">вручную</td>'
            else:
                cells += f'<td>{num(tot[key])}&nbsp;ч<span>{round(100 * tot[key] / s)}%</span></td>'
        rows += f'<tr class="wsdir-{cls}"><th>{year}</th>{cells}</tr>'
    return f'<table class="wsdir-legend"><thead><tr><th></th>{head}</tr></thead><tbody>{rows}</tbody></table>'


CSS = '''  /* wsdir:start — модалка «Помесячно по направлениям» (tools/directions_modal.py) */
  .wsdir-head { display: flex; justify-content: space-between; align-items: center; margin-bottom: 28px; }
  .wsdir-head > div { margin-bottom: 0 !important; }
  .wsdir-btn { font-size: 10px; font-weight: 600; letter-spacing: 0.14em; text-transform: uppercase; color: rgba(230,230,240,0.75); text-decoration: none; padding: 7px 14px; border: 1px solid rgba(255,255,255,0.22); border-radius: 999px; white-space: nowrap; }
  .wsdir-btn:hover { color: #fff; border-color: rgba(255,255,255,0.55); background: rgba(255,255,255,0.05); }
  .wsdir-modal { display: none; position: fixed; inset: 0; z-index: 1000; align-items: center; justify-content: center; }
  .wsdir-modal:target, .wsdir-modal:has(.wsdir-detail:target) { display: flex; }
  .wsdir-col { cursor: pointer; }
  .wsdir-detail { display: none; position: absolute; inset: 0; z-index: 2; align-items: center; justify-content: center; background: rgba(10,10,10,0.82); border-radius: 14px; }
  .wsdir-detail:target { display: flex; }
  .wsdir-dcard { width: min(760px, 90%); max-height: 92%; overflow: auto; background: #0e0e10; border: 1px solid rgba(255,255,255,0.18); border-radius: 12px; padding: 22px 28px 18px; box-shadow: 0 24px 50px rgba(0,0,0,0.6); }
  .wsdir-back { display: inline-block; margin-bottom: 12px; font-size: 10px; font-weight: 600; letter-spacing: 0.14em; text-transform: uppercase; color: rgba(230,230,240,0.7); text-decoration: none; padding: 6px 12px; border: 1px solid rgba(255,255,255,0.22); border-radius: 999px; }
  .wsdir-back:hover { color: #fff; border-color: rgba(255,255,255,0.55); }
  .wsdir-dh { font-size: 18px; font-weight: 700; color: rgba(255,255,255,0.96); }
  .wsdir-dsub { margin: 2px 0 12px; font-size: 11px; color: rgba(230,230,240,0.5); font-variant-numeric: tabular-nums; }
  .wsdir-drow { display: grid; grid-template-columns: 20px minmax(0, 1fr) 58px 140px 72px; align-items: center; gap: 10px; padding: 6px 0; border-bottom: 1px solid rgba(255,255,255,0.07); font-size: 12px; color: rgba(255,255,255,0.88); font-variant-numeric: tabular-nums; }
  .wsdir-dn { color: rgba(230,230,240,0.4); font-size: 11px; }
  .wsdir-dp { overflow: hidden; text-overflow: ellipsis; white-space: nowrap; }
  .wsdir-dg { font-size: 9px; font-weight: 600; letter-spacing: 0.1em; text-align: center; padding: 2px 0; border-radius: 3px; }
  .wsdir-dg-vo { background: rgba(255,255,255,0.92); color: #0a0a0a; }
  .wsdir-dg-dpo { background: rgba(255,255,255,0.42); color: #0a0a0a; }
  .wsdir-dg-teens { border: 1px solid rgba(255,255,255,0.6); color: rgba(255,255,255,0.9); }
  .wsdir-dbar { height: 6px; background: rgba(255,255,255,0.08); }
  .wsdir-dbar i { display: block; height: 100%; background: rgba(255,255,255,0.6); }
  .wsdir-dv { text-align: right; font-weight: 600; }
  .wsdir-dempty { padding: 16px 0; font-size: 12px; color: rgba(230,230,240,0.45); }
  .wsdir-bg { position: absolute; inset: 0; background: rgba(0,0,0,0.88); }
  .wsdir-panel { position: relative; width: min(1240px, 94vw); max-height: 96vh; overflow: auto; background: #0a0a0a; border: 1px solid rgba(255,255,255,0.14); border-radius: 14px; padding: 24px 44px 20px 76px; box-shadow: 0 30px 60px rgba(0,0,0,0.55); }
  .wsdir-close { position: absolute; top: 14px; right: 14px; width: 32px; height: 32px; border-radius: 50%; background: rgba(255,255,255,0.08); display: flex; align-items: center; justify-content: center; font-size: 16px; font-weight: 600; color: rgba(230,230,240,0.75); text-decoration: none; }
  .wsdir-close:hover { color: #fff; }
  .wsdir-h { font-size: 20px; font-weight: 700; color: rgba(255,255,255,0.96); margin-bottom: 4px; }
  .wsdir-sub { font-size: 12px; line-height: 1.5; color: rgba(230,230,240,0.5); margin-bottom: 8px; }
  .wsdir-sub b { font-weight: 600; color: rgba(230,230,240,0.75); }
  .wsdir-kind { display: flex; justify-content: space-between; align-items: baseline; margin: 12px 0 6px; font-size: 11px; font-weight: 600; letter-spacing: 0.18em; text-transform: uppercase; color: rgba(230,230,240,0.6); }
  .wsdir-kind span { letter-spacing: 0.1em; color: rgba(230,230,240,0.4); }
  .wsdir-plot { position: relative; }
  .wsdir-svg { display: block; overflow: visible; }
  .wsdir-col:hover rect:not([fill="transparent"]) { filter: brightness(1.35); }
  .wsdir-col:hover rect[fill="transparent"] { fill: rgba(255,255,255,0.05); }
  .wsdir-tick { position: absolute; right: 100%; padding-right: 8px; transform: translateY(-50%); font-size: 9px; font-variant-numeric: tabular-nums; color: rgba(230,230,240,0.35); white-space: nowrap; }
  .wsdir-months { display: grid; grid-template-columns: repeat(24, 1fr); margin-top: 6px; font-size: 8px; letter-spacing: 0.04em; text-align: center; color: rgba(230,230,240,0.35); }
  .wsdir-years { display: grid; grid-template-columns: 1fr 1fr; margin-top: 4px; font-size: 10px; font-weight: 600; letter-spacing: 0.14em; text-align: center; color: rgba(230,230,240,0.55); }
  .wsdir-legend { width: 100%; margin-top: 8px; border-collapse: collapse; table-layout: fixed; font-size: 12px; font-variant-numeric: tabular-nums; color: rgba(255,255,255,0.88); }
  .wsdir-legend th, .wsdir-legend td { padding: 4px 10px 4px 0; text-align: left; font-weight: 400; vertical-align: baseline; }
  .wsdir-legend thead th { padding-bottom: 6px; border-bottom: 1px solid rgba(255,255,255,0.14); }
  .wsdir-legend thead th:first-child, .wsdir-legend tbody th { width: 76px; }
  .wsdir-legend thead b { font-weight: 700; }
  .wsdir-legend thead small { display: block; margin: 2px 0 0 20px; font-size: 10px; color: rgba(230,230,240,0.4); overflow: hidden; text-overflow: ellipsis; white-space: nowrap; }
  .wsdir-legend tbody th { font-size: 10px; font-weight: 600; letter-spacing: 0.12em; color: rgba(230,230,240,0.55); }
  .wsdir-legend tbody tr { border-bottom: 1px solid rgba(255,255,255,0.07); }
  .wsdir-legend td span { margin-left: 8px; color: rgba(230,230,240,0.45); }
  .wsdir-y1 td { color: rgba(230,230,240,0.6); }
  .wsdir-y2 td { font-weight: 600; }
  .wsdir-legend td.wsdir-na { font-weight: 400; font-size: 11px; color: rgba(230,230,240,0.35); }
  .wsdir-sw { display: inline-block; width: 12px; height: 12px; margin-right: 8px; vertical-align: -1px; }
  .wsdir-sw-vo { background: rgba(255,255,255,0.92); }
  .wsdir-sw-dpo { background: rgba(255,255,255,0.42); }
  .wsdir-sw-teens { background: repeating-linear-gradient(45deg, rgba(255,255,255,0.85) 0 1.5px, rgba(255,255,255,0.12) 1.5px 3.5px); }
  .wsdir-sw-spo { background: rgba(255,255,255,0.18); }
  /* wsdir:end */
'''

for page in sys.argv[1:] or list(PAGES):
    prefix = PAGES[page]
    p = ROOT / 'pages' / page / 'index.html'
    t = p.read_text(encoding='utf-8')
    # убрать результат прошлого запуска
    t = re.sub(r'\n?<!-- wsdir:start -->.*?<!-- wsdir:end -->\n?', '\n', t, flags=re.S)
    t = re.sub(r'  /\* wsdir:start.*?/\* wsdir:end \*/\n', '', t, flags=re.S)
    t = re.sub(r'<div class="wsdir-head">(<div class="ws\d+-title">.*?</div>)<a [^>]*class="wsdir-btn"[^>]*>.*?</a></div>', r'\1', t)

    title_m = re.search(rf'<div class="ws{page}-title">(.*?)</div>', t)
    name = title_m.group(1)
    mid = f'ws{page}-dir'
    t = t.replace(title_m.group(0), f'<div class="wsdir-head">{title_m.group(0)}<a href="#{mid}" class="wsdir-btn">Помесячно по направлениям</a></div>', 1)

    body, cards = [], []
    for kind in ('SELF', 'CLASS'):
        data, progs = load(prefix, kind)
        body.append(f'<div class="wsdir-kind">{kind}<span>часы в месяц</span></div>{chart(page, kind, data)}{legend(data)}')
        cards.append(details(page, kind, data, progs))
    body.append('\n    ' + ''.join(cards))
    modal = (f'<!-- wsdir:start -->\n<div id="{mid}" class="wsdir-modal">\n  <a href="#" class="wsdir-bg" aria-label="Закрыть"></a>\n'
             f'  <div class="wsdir-panel">\n    <a href="#" class="wsdir-close">&#10005;</a>\n'
             f'    <div class="wsdir-h">{name} · часы по направлениям помесячно</div>\n'
             f'    <div class="wsdir-sub">Август 2024 — июль 2026. Столбик — сумма часов за месяц, сегменты — группы; наведите на месяц, чтобы увидеть часы, кликните — топ-10 программ месяца. '
             f'Под графиком — итог учебного года и доля группы. '
             f'<b>ВО</b> — {GROUP_NOTE["vo"]}; <b>ДПО</b> — {GROUP_NOTE["dpo"]}; <b>Uteens</b> — {GROUP_NOTE["teens"]}; <b>СПО</b> — {GROUP_NOTE["spo"]}.</div>\n'
             f'    {"".join(body)}\n  </div>\n</div>\n<!-- wsdir:end -->\n')
    t = t.rstrip('\n') + '\n\n' + modal
    t = t.replace('</style>', CSS + '</style>', 1)
    p.write_text(t, encoding='utf-8', newline='\n')
    print(f'page {page}: ok')
